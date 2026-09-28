from openpyxl import load_workbook
from os import path
from uuid import uuid4, UUID
from datetime import datetime
from pydantic import BaseModel
import json
from datetime import date

# TODO: Add tkinter file picker


class Word(BaseModel):
    wordId: UUID
    chapterId: UUID
    word: str
    translation: str
    targetDate: datetime


class Chapter(BaseModel):
    chapterId: UUID
    name: str
    words: list[Word]


def get_current_chapter():
    if not chapters:
        raise ValueError("No chapter has been added yet.")
    return chapters[-1]


def _to_datetime(d):
    """Return a datetime for a date or datetime input."""
    if d is None:
        return None
    if isinstance(d, datetime):
        return d
    if isinstance(d, date):
        return datetime(d.year, d.month, d.day)
    raise TypeError("Unsupported date type: %r" % type(d))


def _adjust_year_rollover(new_dt: datetime, last_dt: datetime | None):
    """If new_dt's month/day is not after last_dt, advance years until it is.

    This handles spreadsheets that omit the year and expect entries to continue
    into the next calendar year (e.g. 12-12 then 1-1 -> 1-1 of next year).
    """
    if new_dt is None:
        return None
    if last_dt is None:
        return new_dt

    # If the new date is already strictly after last_dt, return it.
    if new_dt > last_dt:
        return new_dt

    # Otherwise, keep adding years until it's after last_dt.
    year = last_dt.year
    try:
        candidate = new_dt.replace(year=year)
    except ValueError:
        # Handle Feb 29 on non-leap year by moving to Feb 28
        candidate = new_dt.replace(year=year, day=28)

    while candidate <= last_dt:
        year += 1
        try:
            candidate = new_dt.replace(year=year)
        except ValueError:
            candidate = new_dt.replace(year=year, day=28)

    return candidate


def generate_chapter(name):
    global current_chapter
    new_chapter = Chapter(chapterId=uuid4(), name=name, words=[])
    chapters.append(new_chapter)
    current_chapter = new_chapter


def generate_word(number: int, word: str, translation: str, target_date: datetime):
    word_id = uuid4()
    chapter = get_current_chapter()
    new_word = Word(
        wordId=word_id,
        chapterId=chapter.chapterId,
        word=word,
        translation=translation,
        targetDate=target_date,
    )
    chapter.words.append(new_word)


def parse_row(row):
    global last_date

    if row == (None, None, None, None, None, None):
        return

    if isinstance(row[1], str):
        # Row is chapter header
        generate_chapter(row[1])
        return

    if row[4]:
        candidate = _to_datetime(row[4])
        last_date = _adjust_year_rollover(candidate, last_date)

    generate_word(
        number=row[1],
        word=row[2],
        translation=row[3],
        target_date=last_date,
    )


def convert_to_json(class_year: int, chapters: list[Chapter]):

    def compute_school_year(chapters: list[Chapter]) -> str:
        """Return a two-year label like '25-26' computed from word target dates.

        Rules:
        - Collect all `targetDate` values from words.
        - If there are dates spanning multiple years, return "YY-YY" using min and max years.
        - If all dates are in the same year, return "YY-(YY+1)" assuming a school year spans two calendar years.
        - If no dates are available, return the current year and next year (fallback).
        """
        years = []
        for chap in chapters:
            for w in chap.words:
                if w.targetDate:
                    years.append(w.targetDate.year)

        if not years:
            y = datetime.now().year
            return f"{y % 100:02d}-{(y + 1) % 100:02d}"

        min_year = min(years)
        max_year = max(years)

        if max_year > min_year:
            return f"{min_year % 100:02d}-{max_year % 100:02d}"
        return f"{min_year % 100:02d}-{(min_year + 1) % 100:02d}"

    learn_year = compute_school_year(chapters)
    return json.dumps(
        {
            "listId": str(uuid4()),
            "year": learn_year,
            "class": class_year,
            "chapters": [chapter.model_dump(mode="json") for chapter in chapters],
        },
        indent=2,
    )



def write_to_file(filename, json_text):
    with open(filename, "w+") as f:
        f.writelines(json_text)


if __name__ == "__main__":
    file = input(
        "Path of file(make sure to use the absolute path, C:/Users/name/Downloads/file):\n"
    )
    while True:
        if path.exists(file):
            x = path.splitext(file)
            if x[1] != ".xlsx":
                print("File is not correct Excel type(.xlsx), please try again:")
            else:
                break
        else:
            print("File wasn't found, please try again:")
        file = input()

    sheetname = "Sheet1"

    wb = load_workbook(file, read_only=True, data_only=True)
    if wb.sheetnames[0] != "Sheet1":
        sheetname = input(
            f"The first sheet in this file isn't default(Sheet1), these are all the sheets in this file:\n{', '.join(wb.sheetnames)} \nplease provide the correct name:\n"
        )
    ws = wb[sheetname]

    class_year = input("What year is the class(1, 2, 3, 4, 5, 6):\n")

    chapters = []
    global current_chapter
    current_chapter = None

    last_date = None

    for row in ws.iter_rows(min_row=2, values_only=True):
        parse_row(row)

    converted = convert_to_json(class_year, chapters)

    new_file_path = path.splitext(file)[0] + ".json"
    write_to_file(new_file_path, converted)
